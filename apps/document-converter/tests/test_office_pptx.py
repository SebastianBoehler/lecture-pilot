from pptx import Presentation
from pptx.util import Inches

from lecturepilot_converter.office_pptx import pptx_supplemental_blocks


def test_grouped_shapes_preserve_nested_shape_and_text_links(tmp_path):
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    group = slide.shapes.add_group_shape()
    nested = group.shapes.add_group_shape()
    textbox = nested.shapes.add_textbox(0, 0, Inches(2), Inches(1))
    textbox.text = "Course evidence"
    textbox.click_action.hyperlink.address = "https://example.edu/course"
    textbox.text_frame.paragraphs[0].runs[0].hyperlink.address = "https://example.edu/evidence"
    group.shapes.add_group_shape()
    source = tmp_path / "grouped.pptx"
    presentation.save(source)

    blocks = pptx_supplemental_blocks(source)

    assert [(block["url"], block["text"], block["locator"]) for block in blocks] == [
        ("https://example.edu/course", "Course evidence", {"slide": 1}),
        ("https://example.edu/evidence", "Course evidence", {"slide": 1}),
    ]
